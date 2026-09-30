"""Regenerate the editable Chakra deck with python-pptx."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, DARK, INK, MUTED, ORANGE, RULE = 'F6F3EE', '15100D', '15100D', '75665B', 'EC7A3A', 'D4C9BD'
TX = 'https://testnet.arcscan.app/tx/0xee7bc19a990ce6691a68e9b387585baee13edc846cbf3a43551ab3dd7cfcda6c'
DEMO = 'https://youtu.be/12hvXfd5rsE'

def box(s, x, y, w, h, color):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh

def text(s, value, x, y, w, h, size=18, color=INK, bold=False, font='Arial', url=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.space_after = Pt(10)
        p.font.name = font; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        if url:
            for run in p.runs:
                run.hyperlink.address = url
                run.font.color.rgb = RGBColor.from_string(color)
                run.font.underline = False
    return sh

def slide(n, label, title, subtitle='', dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(DARK if dark else BG)
    fg = BG if dark else INK
    text(s, f'CHAKRA  /  {label.upper()}', .7, .45, 11.7, .3, 11, ORANGE, True)
    text(s, title, .7, 1.05, 11.8, 1.35, 38, fg, True)
    if subtitle: text(s, subtitle, .72, 2.45, 11.65, .7, 18, 'B8AFA6' if dark else MUTED)
    box(s, .7, 6.91, 11.9, .01, '4A3C31' if dark else RULE)
    text(s, 'CIRCLE DEVELOPER GRANT  /  SEPTEMBER 2026', .7, 7.06, 10, .2, 9, 'B8AFA6' if dark else MUTED)
    text(s, f'{n:02d} / 11', 11.8, 7.03, .8, .25, 10, ORANGE)
    return s

def column(s, x, kicker, title, body, w=3.65, dark=False):
    box(s, x, 3.35, w, .025, ORANGE)
    text(s, kicker, x, 3.6, w, .35, 11, ORANGE, True)
    text(s, title, x, 4.13, w, .7, 23, BG if dark else INK, True)
    text(s, body, x, 5.0, w, 1.65, 16, 'B8AFA6' if dark else MUTED)

s = slide(1, 'Stablecoin FX on Arc', 'Stablecoin swaps.\nOne execution interface.', 'USDC / EURC routing for wallets and payments apps.', True)
text(s, 'Chakra', .7, 3.45, 8, 1, 64, ORANGE, True)
text(s, 'Compare supported venues. Build the transaction.\nThe user authorizes in their wallet.', .75, 4.75, 10, 1.0, 23, BG)
text(s, 'Open source  ·  Arc Testnet proof  ·  Proposed Circle integration', .75, 6.2, 11.5, .4, 14, 'B8AFA6')
s.notes_slide.notes_text_frame.text = 'Existing testnet execution is separate from proposed Circle Wallets, Gas Station and CCTP work.'

s = slide(2, 'Problem', 'Every venue adds integration work.', 'A USDC/EURC swap spans liquidity selection, execution and wallet authorization.')
column(s,.7,'01 / LIQUIDITY','Different pool math','Apps must compare supported venues and handle missing or thin routes.')
column(s,4.82,'02 / EXECUTION','Different interfaces','A quote must become executable calldata with approvals and a deadline.')
column(s,8.94,'03 / WALLET','Authorization + gas','Embedded wallets need a reliable payload and bounded sponsorship eligibility.')

s = slide(3, 'Product', 'Quote → build → authorize → settle', 'One API and web app, with venue-aware routing and wallet-ready transaction data.')
steps=[('QUOTE','Compare supported pools'),('BUILD','Unsigned calldata'),('AUTHORIZE','User wallet approval'),('SETTLE','Atomic on-chain route')]
for i,(a,b) in enumerate(steps):
    x=.7+i*3.05
    box(s,x,3.65,2.72,1.6,'FFFDF9'); text(s,a,x+.18,3.9,2.4,.35,15,ORANGE,True); text(s,b,x+.18,4.45,2.35,.6,18)
text(s,'Factory-gated venues  /  Permit2 pull  /  Fail closed when no route exists',.7,5.75,11.8,.65,20,MUTED)

s = slide(4, 'Executed proof', '1 USDC → 1.629188 EURC', 'Recorded Arc Testnet wallet execution · 5 September 2026 · Presto single route', True)
text(s,'60563600',.7,3.5,5.7,1,54,ORANGE,True)
text(s,'BLOCK / TESTNET CHAIN 5042002',.75,4.65,5.7,.45,12,'B8AFA6')
text(s,'MetaMask authorization\nSuccessful splitSwap entry point\nSingle route; no mined split claimed',7.0,3.65,5.0,1.9,21,BG)
text(s,'Open transaction receipt ↗',.75,5.65,6,.5,18,ORANGE,url=TX)
text(s,'Open recorded demo ↗',7,5.65,5.1,.5,18,ORANGE,url=DEMO)
s.notes_slide.notes_text_frame.text=f'Receipt: {TX}\nDemo: {DEMO}\nReceipt details recovered from existing testing records. Explorer API returned HTTP 403 during this refresh; do not claim fresh receipt verification.'

s = slide(5, 'Activity baseline', 'Execution evidence before scale.', 'Public stats API fetched 30 September 2026. Testnet activity does not establish customer acquisition.')
for x,num,label in [(.7,'5','INDEXED SWAPS'),(4.82,'5 USDC','NOTIONAL'),(8.94,'2','UNIQUE ADDRESSES')]:
    text(s,num,x,3.6,3.65,.9,44,INK,True); text(s,label,x,4.65,3.65,.4,12,ORANGE,True)
text(s,'0 split swaps. Index freshness lag: about 44 hours.\nNo verified customer count, organic volume or revenue claimed.',.7,5.55,11.9,.9,18,MUTED)
s.notes_slide.notes_text_frame.text='Source: https://chakra-api-0a5i.onrender.com/api/v1/stats?range=all\nSaved: docs/evidence/grant-testnet/stats-2026-09-30.json\nfreshness_secs=157662. Historical 5 September snapshot was 3 swaps / 3 USDC / 1 trader.'

s = slide(6, 'Funding boundary', 'A working router. A new Circle path.', 'The request funds incremental integration, verification and documentation.')
column(s,.7,'EXISTING BASELINE','Arc Testnet','Router, Solidity execution, quote/build API, SDK source and external-wallet receipt.')
column(s,4.82,'CIRCLE REQUEST','Wallet + gas path','New SCA adapter, mined smart-account swap, separate sponsored swap and CCTP handoff.')
column(s,8.94,'SELF-FUNDED','Mainnet proof','Mainnet deployment and external-wallet proof are prerequisites, excluded from this ask.')

s = slide(7, 'Milestone plan', 'Four jobs. Twenty sequential weeks.', '24,000 USDC requested · one engineer · evidence required at every milestone')
rows=[('01','Circle SDK adapter','4 weeks','4,800','npm release, examples, authorization mapping'),('02','Circle SCA swap','7 weeks','8,400','Mined Testnet swap + mainnet readiness'),('03','Gas Station sponsorship','4 weeks','4,800','Separate mined swap + matching sponsorship record'),('04','CCTP V2 inbound','5 weeks','6,000','Burn/receive receipts + post-receive Chakra result')]
for i,(n,title,weeks,cost,proof) in enumerate(rows):
    y=3.35+i*.78
    box(s,.7,y,11.9,.015,RULE)
    text(s,n,.7,y+.15,.6,.4,15,ORANGE,True)
    text(s,title,1.42,y+.12,4.45,.4,20,INK,True)
    text(s,proof,1.42,y+.5,7.6,.25,11,MUTED)
    text(s,weeks,9.4,y+.16,1.35,.4,16,MUTED)
    text(s,cost,11.05,y+.12,1.5,.45,21,INK,True)

s = slide(8, 'Reviewer checks', 'Acceptance lives in the receipts.', 'Configuration and simulations support the proof; mined execution has its own gate.')
column(s,.7,'WALLETS','One mined SCA swap','Testnet receipt linked to quote/build and Circle authorization; mainnet configuration with a live route.')
column(s,4.82,'GAS STATION','A different receipt','Separate Testnet transaction matched to a redacted sponsorship record; bounded policy and failure checks.')
column(s,8.94,'CCTP V2','One linked transfer','Source burn and Arc receive; then Chakra quote and swap, or a dated no-route result.')

s = slide(9, 'Distribution', 'Start with one wallet integrator.', 'A proposed first-adopter target, followed by a public and reproducible integration record.')
column(s,.7,'ADOPTION','First live swap','Publish a named integrator and its receipt if achieved. Track time to first swap and dated npm downloads.')
column(s,4.82,'OPERATIONS','Observable API','Report rate limits, 429s and stated-interval uptime. Publish dated swap count, volume and pair mix.')
column(s,8.94,'REVENUE PLAN','2 bps on mainnet','Proposed 0.02% fee on successful swaps. 1M USDC routed would yield 200 USDC before costs; current Testnet fee is zero.')

s = slide(10, 'Founder', 'Built by an engineer who ships.', 'Kyler Nguyen · technical founder · contract, routing, API and integration work')
text(s,'Solidity + Rust + TypeScript',.7,3.5,11.8,.6,28,INK,True)
text(s,'Public contributions across Chakra, Kamui, MGK Protocol and Oreka:\nDEX routing, VRF, perpetuals and prediction-market infrastructure.',.7,4.42,11.6,1.1,22,MUTED)
text(s,'View founder GitHub ↗',.7,5.95,5.7,.45,18,ORANGE,url='https://github.com/MrSufferer')
text(s,'One engineer. Sequential scope.',7,5.95,5.5,.45,18,INK)
s.notes_slide.notes_text_frame.text='GitHub API contributor records verify MrSufferer contributions. Do not imply sole authorship, commercial adoption, partnerships or awards. https://github.com/mangekyou-labs/chakra https://github.com/mangekyou-labs/kamui https://github.com/mangekyou-labs/mgk-protocol https://github.com/mangekyou-labs/oreka'

s = slide(11, 'The request', '24,000 USDC', 'Bring the existing Chakra route into a verifiable Circle wallet and sponsorship flow.', True)
text(s,'800 engineer hours × 30 USDC/hour',.7,3.65,11.8,.65,28,BG,True)
text(s,'20 weeks · four milestones · no separate expense or markup\nExisting router and self-funded mainnet proof excluded',.7,4.6,11.8,1.1,21,'B8AFA6')
text(s,'Watch demo ↗',.7,6.15,3.6,.35,16,ORANGE,url=DEMO)
text(s,'Explore product ↗',4.82,6.15,3.65,.35,16,ORANGE,url='https://chakra-ag.vercel.app')
text(s,'Inspect source ↗',8.94,6.15,3.65,.35,16,ORANGE,url='https://github.com/mangekyou-labs/chakra')
prs.core_properties.author='Kyler Nguyen'
prs.core_properties.title='Chakra — Circle Developer Grant'
prs.core_properties.subject='24,000 USDC / four milestones / Arc stablecoin FX'
output = BytesIO()
prs.save(output)
# LibreOffice uses theme hyperlink colors even with explicit run colors.
with ZipFile(output) as source, ZipFile(ROOT/'Chakra-Circle-Grant-Pitch-Deck.pptx', 'w', ZIP_DEFLATED) as target:
    for item in source.infolist():
        data = source.read(item.filename)
        if item.filename.startswith('ppt/theme/') and item.filename.endswith('.xml'):
            root = etree.fromstring(data)
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for node in root.xpath('//a:clrScheme/a:hlink | //a:clrScheme/a:folHlink', namespaces=ns):
                for child in list(node): node.remove(child)
                etree.SubElement(node, '{'+ns['a']+'}srgbClr', val=ORANGE)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        target.writestr(item, data)
print('Generated 11 editable slides')
